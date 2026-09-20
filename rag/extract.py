"""Pulls raw text out of the source documents, keeping page/slide-level provenance."""
from pathlib import Path
import re

import fitz  # pymupdf
from pptx import Presentation


def extract_pdf(path: Path):
    """Returns a list of {text, loc, source} dicts, one per PDF page."""
    doc = fitz.open(path)
    out = []
    for i, page in enumerate(doc, start=1):
        text = page.get_text().strip()
        if text:
            out.append({"text": text, "loc": f"page {i}", "source": path.name})
    return out


def extract_pptx(path: Path):
    """Returns a list of {text, loc, source} dicts, one per slide (title + body text)."""
    prs = Presentation(path)
    out = []
    for i, slide in enumerate(prs.slides, start=1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text.strip())
            if shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        text = "\n".join(parts).strip()
        if text:
            out.append({"text": text, "loc": f"slide {i}", "source": path.name})
    return out


def extract_py(path: Path):
    """Splits a python source file into top-level function/class blocks."""
    src = path.read_text(encoding="utf-8")
    pattern = re.compile(r"^(def |class )", re.MULTILINE)
    starts = [m.start() for m in pattern.finditer(src)]
    if not starts:
        return [{"text": src, "loc": "whole file", "source": path.name}]

    out = []
    header = src[: starts[0]].strip()
    if header:
        out.append({"text": header, "loc": "module header", "source": path.name})

    for idx, start in enumerate(starts):
        end = starts[idx + 1] if idx + 1 < len(starts) else len(src)
        block = src[start:end].strip()
        name_match = re.match(r"(def|class)\s+(\w+)", block)
        name = name_match.group(2) if name_match else f"block_{idx}"
        out.append({"text": block, "loc": f"function/class {name}", "source": path.name})
    return out


def extract_file(path: Path):
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_pdf(path)
    if suffix == ".pptx":
        return extract_pptx(path)
    if suffix == ".py":
        return extract_py(path)
    raise ValueError(f"Unsupported file type: {path}")


def extract_corpus(raw_dir: Path):
    """Walks raw_dir and extracts every supported file into page/slide/block units."""
    units = []
    for path in sorted(raw_dir.iterdir()):
        if path.suffix.lower() in (".pdf", ".pptx", ".py"):
            units.extend(extract_file(path))
    return units
